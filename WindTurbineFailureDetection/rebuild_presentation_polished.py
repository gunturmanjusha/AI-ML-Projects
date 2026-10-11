import json,os
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
r=json.load(open('outputs/results.json'));p=Presentation('AIML_Project_1_Low_Code_Presentation_Template.pptx');W,H=p.slide_width,p.slide_height
NAVY=RGBColor(15,55,80);TEAL=RGBColor(0,139,143);PALE=RGBColor(231,244,244);INK=RGBColor(35,49,58);WHITE=RGBColor(255,255,255)
def clear(s):
 for sh in list(s.shapes):
  s.shapes._spTree.remove(sh._element)
def head(s,title,section):
 clear(s);bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,W,Inches(.75));bar.fill.solid();bar.fill.fore_color.rgb=NAVY;bar.line.fill.background()
 t=s.shapes.add_textbox(Inches(.55),Inches(.12),Inches(10.8),Inches(.4)).text_frame;t.text=title;t.paragraphs[0].font.size=Pt(25);t.paragraphs[0].font.bold=True;t.paragraphs[0].font.color.rgb=WHITE
 f=s.shapes.add_textbox(Inches(.55),Inches(7.0),Inches(12),Inches(.25)).text_frame;f.text=section+'  |  Wind Turbine Failure Detection';f.paragraphs[0].font.size=Pt(9);f.paragraphs[0].font.color.rgb=RGBColor(100,110,115)
def text(s,content,x,y,w,h,size=18,color=INK,bold=False):
 tf=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame;tf.clear()
 for i,line in enumerate(content):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.size=Pt(size);q.font.color.rgb=color;q.font.bold=bold if i==0 else False;q.space_after=Pt(10)
 return tf
def pic(s,path,x,y,w,h=None):
 if os.path.exists(path):s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w),height=Inches(h) if h else None)
def tbl(s,data,x,y,w,h):
 t=s.shapes.add_table(len(data),len(data[0]),Inches(x),Inches(y),Inches(w),Inches(h)).table
 for i,row in enumerate(data):
  for j,v in enumerate(row):
   c=t.cell(i,j);c.text=str(v);c.fill.solid();c.fill.fore_color.rgb=NAVY if i==0 else (PALE if i%2 else WHITE)
   for q in c.text_frame.paragraphs:q.font.size=Pt(11);q.font.color.rgb=WHITE if i==0 else INK;q.alignment=PP_ALIGN.CENTER;q.font.bold=i==0
# Replace content of all template slides
for s in p.slides: clear(s)
# title
s=p.slides[0];head(s,'Wind Turbine Failure Detection','Great Learning | Low-Code ML Assignment');text(s,['Aeolus Renewables','Drivetrain condition classification from 10-minute SCADA data','Validated timestamp-grouped chronological evaluation'],.8,1.5,8,2,25,TEAL,True);pic(s,'outputs/charts/eda_target_turbines.png',8,1.3,4.6,3)
s=p.slides[1];head(s,'Contents / Agenda','Main Presentation');text(s,['1  Executive summary and business problem','2  Exploratory analysis and preprocessing','3  Baseline models, tuning and final selection','4  Final performance, feature interpretation and recommendations','Appendix: data background, baseline table, tuning details and limitations'],.9,1.35,11,4.8,22)
s=p.slides[2];head(s,'Executive Summary','Main Presentation');text(s,['Tuned ANN selected: 92.49% validation recall','Test recall: 86.40% — 1,245 of 1,441 recorded fault states identified','196 missed states remain the principal risk; 1,271 false positives require triage','Deploy as current-condition decision support, not as a proven future-failure predictor'],.7,1.25,6.1,4.9,20);pic(s,'outputs/charts/final_test_confusion_matrix.png',7.1,1.35,5.0,4.1)
s=p.slides[3];head(s,'Business Problem Overview and Solution Approach','Main Presentation');text(s,['Business need','Reduce manual SCADA review by prioritising turbines with recorded drivetrain fault signatures.'],.7,1.2,5.7,1.3,19,TEAL,True);text(s,['Method','Timestamp-grouped 70/15/15 chronology; training-only imputation/scaling; five classifiers; tune ANN and XGBoost; select on validation recall; evaluate once on test.'],.7,3.0,11.4,2.1,19)
s=p.slides[4];head(s,'EDA Results — Fault Pattern and Sensor Evidence','Main Presentation');text(s,['131,760 records from 15 turbines; overall recorded-failure prevalence 3.00%.','Fault-state distributions shift in gearbox temperature, drivetrain vibration and oil-particle readings.','Use signal combinations rather than a single fixed threshold.'],.65,1.0,12,1.1,15);pic(s,'outputs/charts/eda_failure_signals.png',.7,2.2,6.0,3.9);pic(s,'outputs/charts/eda_correlation.png',6.85,2.2,5.7,3.9)
s=p.slides[5];head(s,'Data Preprocessing and Leakage Prevention','Main Presentation');text(s,['Data quality','0 duplicates; 645 gearbox-oil, 690 generator-bearing and 636 oil-pressure missing readings.','Treatment','Median imputation fitted on training rows only; StandardScaler fitted on training rows only for ANN.','Features','Calendar cycles, power utilisation, gearbox and generator temperature rise.','Split','Timestamp-grouped chronological partitions prevent the same SCADA timestamp appearing in two splits.'],.75,1.1,11.5,5.5,18)
s=p.slides[6];head(s,'Model Performance Summary and Final Model Selection','Main Presentation');tbl(s,[['Tuned model','Validation recall','F1','Precision','Decision'],['ANN','92.49%','53.54%','37.67%','Selected'],['XGBoost','69.10%','69.34%','69.57%','Alternative']],.75,1.4,11.5,1.5);text(s,['Selection rule: highest validation recall; F1 then precision only if recall ties.','ANN wins on recall; XGBoost is more precise and has better F1, but misses more recorded faults.','Test data did not influence tuning, selection or threshold choice.'],.8,3.35,11.2,2.5,18)
s=p.slides[7];head(s,'Actionable Business Recommendations','Main Presentation');text(s,['1. Rank persistent high-risk alerts for operations-centre review.','2. Check vibration, thermal, oil and maintenance context before dispatch.','3. Send recurring or multi-signal alerts to planned inspection queues.','4. Monitor missed states, false-alert workload, sensor drift and downtime outcomes.','5. Validate prospectively before claiming lead time or financial savings.'],.8,1.15,11.4,5.2,20)
s=p.slides[8];head(s,'APPENDIX — Data Background and Contents','Appendix');text(s,['SCADA panel data: 8,784 ten-minute timestamps across 15 turbines.','Input groups: weather, production, rotational speed, temperatures, vibration spectra, oil condition, operating history and maintenance context.','Target: failure indicates the condition recorded at the timestamp.'],.8,1.3,11.3,4,20)
s=p.slides[9];head(s,'Model Building — Baseline Validation Comparison','Appendix');data=[['Model','Accuracy','Precision','Recall','F1']]+[[q['model'],f"{q['accuracy']:.1%}",f"{q['precision']:.1%}",f"{q['recall']:.1%}",f"{q['f1']:.1%}"] for q in r['baseline_validation']];tbl(s,data,.6,1.45,12.1,2.8);pic(s,'outputs/charts/baseline_recall.png',2.0,4.55,9.1,1.8)
s=p.slides[10];head(s,'Model Improvement — Hyperparameter Tuning','Appendix');text(s,['XGBoost search: 6 randomized configurations; two timestamp-grouped expanding folds; recall optimization.','Ranges: 80–200 trees, depth 3–6, learning rate 0.02–0.10, row/column sampling 0.70–1.00, min-child weight 1–5.','Best XGBoost CV recall 48.45%; holdout validation recall 69.10%.','ANN: 64/24 layers, BatchNorm, dropout 0.30/0.15, Adam and early stopping.'],.7,1.05,11.6,4.4,18)
s=p.slides[11];head(s,'Final Model Performance — Untouched Test Set','Appendix');text(s,['Tuned ANN','Accuracy 92.58% | Precision 49.48% | Recall 86.40% | F1 62.93%','TN 17,058 | FP 1,271 | FN 196 | TP 1,245','Interpretation: false negatives require monitoring; false positives require remote review and alert-persistence rules.'],.75,1.25,6.2,3.6,18,TEAL,True);pic(s,'outputs/charts/final_test_confusion_matrix.png',7.15,1.25,4.8,4.1)
s=p.slides[12];head(s,'Feature Interpretation and Deployment Limitations','Appendix');text(s,['Permutation importance measures validation-recall reduction after shuffling an input.','It identifies associated condition signals, not causal mechanisms.','Use top signals to focus technician checks; retrain and monitor drift as operations change.'],.7,1.05,5.8,3,18);pic(s,'outputs/charts/feature_importance.png',6.5,1.1,5.7,4.7)
# Add 3 appendix support slides
for title,lines,img in [('Timestamp-Grouped Split Verification',['Train: 92,220 rows | 6,148 timestamps | 1.60% faults','Validation: 19,770 rows | 1,318 timestamps | 5.26% faults','Test: 19,770 rows | 1,318 timestamps | 7.29% faults','No timestamp occurs in more than one partition.'],None),('EDA Supporting View',['Class distribution and turbine-level failure patterns frame the imbalanced classification task.','Failure prevalence differs over time; accuracy alone is insufficient.'],'outputs/charts/eda_target_turbines.png'),('Operational Use Case',['Operations centre reviews alert evidence remotely.','Technician dispatch follows persistence or corroborating signals.','Feedback from work orders supports recalibration and governance.'],None)]:
 s=p.slides.add_slide(p.slide_layouts[6]);head(s,title,'Appendix');text(s,lines,.75,1.25,5.8 if img else 11.4,4.7,20); 
 if img:pic(s,img,6.8,1.35,5.4,3.8)
p.save('outputs/AIML_Project_1_Low_Code_Final_Presentation.pptx');print(len(p.slides))
