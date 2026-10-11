"""Final layout-only correction for the existing 18-slide final deck."""
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

PATH = "outputs/AIML_Project_1_Low_Code_Final_Presentation.pptx"
NAVY = RGBColor(12, 58, 82)
TEAL = RGBColor(0, 132, 137)
PALE = RGBColor(235, 246, 246)
INK = RGBColor(38, 54, 63)
WHITE = RGBColor(255, 255, 255)


def title(slide):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.top < Inches(0.7) and shape.text.strip():
            return shape.text.strip()
    return ""


def find_slide(deck, prefix):
    return next(slide for slide in deck.slides if title(slide).startswith(prefix))


def move_after(deck, moving_prefix, after_prefix):
    source_index = next(i for i, slide in enumerate(deck.slides) if title(slide).startswith(moving_prefix))
    source = deck.slides._sldIdLst[source_index]
    deck.slides._sldIdLst.remove(source)
    target_index = next(i for i, slide in enumerate(deck.slides) if title(slide).startswith(after_prefix))
    deck.slides._sldIdLst.insert(target_index + 1, source)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def add_text(slide, x, y, w, h, text, size, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    return box


def info_card(slide, y, heading, body):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.80), Inches(y), Inches(3.60), Inches(1.58))
    card.fill.solid(); card.fill.fore_color.rgb = PALE
    card.line.color.rgb = TEAL
    add_text(slide, 6.00, y + 0.16, 3.10, 0.22, heading, 11, TEAL, True)
    body_box = add_text(slide, 6.00, y + 0.45, 3.12, 0.96, body, 8.5)
    for paragraph in body_box.text_frame.paragraphs:
        paragraph.space_after = Pt(2)


def update_numbers(deck):
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text
            if re.fullmatch(r"Wind Turbine Failure Detection\s+•\s+\d{2}", text.strip()):
                shape.text_frame.paragraphs[0].text = f"Wind Turbine Failure Detection  •  {number:02d}"
                shape.text_frame.paragraphs[0].font.size = Pt(7)
            elif re.fullmatch(r"(?:MAIN PRESENTATION|APPENDIX)\s+•\s+\d{2}", text.strip()):
                section = "MAIN PRESENTATION" if number <= 9 else "APPENDIX"
                shape.text_frame.paragraphs[0].text = f"{section} • {number:02d}"
                shape.text_frame.paragraphs[0].font.size = Pt(6.5)
                shape.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def main():
    deck = Presentation(PATH)

    # Slide 12: retain chart and rebuild only its right explanatory area.
    tuning = find_slide(deck, "Model Improvement — Hyperparameter Tuning")
    for shape in list(tuning.shapes):
        if shape.left >= Inches(5.70) and shape.top >= Inches(1.0):
            remove(shape)
    info_card(
        tuning, 1.03, "XGBoost Tuning",
        "Six configurations; two timestamp-grouped expanding folds; recall objective.\nSelected: 160 trees, depth 3, learning rate 0.07, subsample 0.85, colsample 1.00, min-child weight 3.\nBest CV recall 48.45%; holdout validation recall 69.10%.",
    )
    info_card(
        tuning, 2.86, "ANN Tuning",
        "Five fixed-seed configurations: baseline, prior 64/24, wider L2, higher-dropout L2, compact regularised. Training-only preprocessing, class weighting and early stopping.\nSelected prior 64/24 ANN: 92.49% recall; best new ANN: 92.40%.",
    )
    # The original right-side footer was part of the replaced explanatory area.
    add_text(tuning, 7.0, 5.28, 2.6, 0.18, "Wind Turbine Failure Detection  •  13", 7)

    # Slide 17: a concise title fits the header without touching the heatmap.
    multivariate = find_slide(deck, "EDA Results — Multivariate")
    for shape in multivariate.shapes:
        if getattr(shape, "has_text_frame", False) and shape.top < Inches(0.7) and shape.text.startswith("EDA Results"):
            shape.text_frame.paragraphs[0].text = "EDA Results — Multivariate Relationships"
            shape.text_frame.paragraphs[0].font.size = Pt(15)
            shape.width = Inches(7.3)
            break

    # Slide 18: exact requested title, retaining its table and selection rationale.
    comparison = find_slide(deck, "ANN and XGBoost Tuning Comparison")
    for shape in comparison.shapes:
        if getattr(shape, "has_text_frame", False) and shape.top < Inches(0.7) and shape.text.startswith("ANN and XGBoost Tuning Comparison"):
            shape.text_frame.paragraphs[0].text = "ANN and XGBoost Tuning Comparison"
            shape.text_frame.paragraphs[0].font.size = Pt(15)
            shape.width = Inches(7.3)
            break

    # Reorder so EDA and tuning appendices remain coherent.
    move_after(deck, "EDA Results — Multivariate", "EDA Results — Sensor Evidence")
    move_after(deck, "ANN and XGBoost Tuning Comparison", "Model Improvement — Hyperparameter Tuning")
    update_numbers(deck)
    deck.save(PATH)
    print("saved", len(deck.slides), "slides")
    for number, slide in enumerate(deck.slides, 1):
        print(number, title(slide))


if __name__ == "__main__":
    main()
