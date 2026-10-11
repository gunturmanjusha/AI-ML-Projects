from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
p=Presentation('outputs/AIML_Project_1_Low_Code_Final_Presentation.pptx');W,H=p.slide_width,p.slide_height
NAVY=RGBColor(12,58,82);TEAL=RGBColor(0,132,137);INK=RGBColor(38,54,63);PALE=RGBColor(235,246,246);WHITE=RGBColor(255,255,255)
def box(s,x,y,w,h,txt,size=16,color=INK,bold=False,fill=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.text=txt
 for q in tf.paragraphs:q.font.size=Pt(size);q.font.color.rgb=color;q.font.bold=bold
 if fill:sh.fill.solid();sh.fill.fore_color.rgb=fill
 return sh
def header(s,title,num,section='MAIN PRESENTATION'):
 b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,W,Inches(.62));b.fill.solid();b.fill.fore_color.rgb=NAVY;b.line.fill.background()
 box(s,.55,.11,9,.35,title,25,WHITE,True);box(s,10.5,.13,2,.28,section+'  •  '+str(num).zfill(2),9,WHITE)
# title date
box(p.slides[0],.75,6.45,4,.28,'Submission date: October 11, 2026',12,INK,True)
# strengthen recs slide 9
box(p.slides[8],.75,5.8,11.6,.75,'Evidence base: 196 missed recorded fault states justify recall-first selection; 1,271 false-positive records justify remote triage before a technician visit. Track alert persistence, missed states, false-alert workload, work-order outcomes and data drift during a pilot.',14,INK,False,PALE)
# add multivariate EDA slide
s=p.slides.add_slide(p.slide_layouts[6]);header(s,'EDA Results — Multivariate Sensor Relationships',17)
box(s,.65,.85,5.35,.85,'Correlation heatmap uses actual SCADA values across thermal, vibration, oil and operating variables.',15,INK)
s.shapes.add_picture('outputs/charts/eda_multivariate_correlation.png',Inches(.6),Inches(1.55),width=Inches(6.1))
box(s,7.0,1.55,5.3,3.9,'What the data shows\n• Drivetrain vibration has the strongest positive correlation with recorded fault condition (0.54).\n• Oil-particle count is also materially positive (0.47); bearing-frequency vibration is 0.23.\n• Oil pressure is negatively related (-0.17); gearbox bearing temperature is weaker (0.10).\n\nBusiness implication\nNo single sensor separates every state. The multi-channel model combines complementary signals and reduces dependence on rigid individual thresholds.',16,INK)
# add tuning comparison
s=p.slides.add_slide(p.slide_layouts[6]);header(s,'Tuning Comparison — Recall First, Precision Trade-off',18,'APPENDIX')
box(s,.65,.85,12,.55,'Both candidates were chosen from baseline validation evidence. Selection used holdout validation recall only; test results were excluded.',16,INK)
data=[['Model','Baseline P / R / F1','Tuned P / R / F1','Tuning method'],['ANN','50.28% / 68.05% / 57.83%','37.67% / 92.49% / 53.54%','64/24 units; BatchNorm; dropout 0.30/0.15; Adam; early stopping'],['XGBoost','71.01% / 60.83% / 65.53%','69.57% / 69.10% / 69.34%','6 randomized configurations; two expanding timestamp-grouped folds']]
t=s.shapes.add_table(3,4,Inches(.65),Inches(1.65),Inches(12.0),Inches(2.0)).table
for i,row in enumerate(data):
 for j,v in enumerate(row):
  c=t.cell(i,j);c.text=v;c.fill.solid();c.fill.fore_color.rgb=NAVY if i==0 else (PALE if i%2 else WHITE)
  for q in c.text_frame.paragraphs:q.font.size=Pt(11 if i else 12);q.font.bold=i==0;q.font.color.rgb=WHITE if i==0 else INK;q.alignment=PP_ALIGN.CENTER
box(s,.8,4.15,11.5,1.65,'Why ANN is selected\nANN has the highest validation recall (92.49%), so it detects more recorded fault states under the declared business priority. XGBoost has higher precision and F1, making it a credible lower-alert-burden alternative. This trade-off is documented rather than hidden.',17,INK,False,PALE)
p.save('outputs/AIML_Project_1_Low_Code_Final_Presentation.pptx');print('slides',len(p.slides))

