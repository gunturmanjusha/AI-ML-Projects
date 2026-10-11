from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
p=Presentation('outputs/AIML_Project_1_Low_Code_Final_Presentation.pptx');W,H=p.slide_width,p.slide_height
NAVY=RGBColor(12,58,82);TEAL=RGBColor(0,132,137);INK=RGBColor(38,54,63);PALE=RGBColor(235,246,246);WHITE=RGBColor(255,255,255)
def set_text(sh,txt,size):
 sh.text_frame.text=txt
 for q in sh.text_frame.paragraphs:q.font.size=Pt(size)
def find(slide,prefix):
 return next(x for x in slide.shapes if hasattr(x,'text') and x.text.startswith(prefix))
# Slide 9: accurate observation-level language
sh=find(p.slides[8],'Evidence base:')
set_text(sh,'Evidence base: 196 missed SCADA fault observations support recall-first selection. The 1,271 false-positive SCADA observations are not 1,271 technician visits or incidents; remote triage and alert persistence should determine dispatch. Pilot KPIs: missed observations, alert workload, work-order confirmation, downtime and drift.',12)
# Slide 12: actual methods only
set_text(find(p.slides[11],'6 sampled'), 'Six randomized XGBoost configurations; two timestamp-grouped expanding-window folds; recall objective.',13)
set_text(find(p.slides[11],'Best CV'), 'Selected XGBoost: 160 trees, depth 3, learning rate 0.07, subsample 0.85, colsample 1.00, min-child weight 3. Best CV recall 48.45%; holdout validation recall 69.10%.',12)
set_text(find(p.slides[11],'64/24'), 'ANN configuration selection: the recorded run compared the baseline ANN with this regularised configuration; no multi-configuration ANN search was recorded. Final ANN: 64/24 units, BatchNorm, dropout 0.30/0.15, Adam, early stopping. Validation recall 92.49%.',12)
# full rebuild targeted slide
def clear(s):
 for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
def tx(s,x,y,w,h,txt,size=16,color=INK,bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.text=txt
 for q in tf.paragraphs:q.font.size=Pt(size);q.font.color.rgb=color;q.font.bold=bold
 return sh
def hdr(s,title,num,sec):
 b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,W,Inches(.62));b.fill.solid();b.fill.fore_color.rgb=NAVY;b.line.fill.background();tx(s,.55,.10,9,.35,title,23,WHITE,True);tx(s,10.3,.14,2.2,.25,sec+' • '+str(num).zfill(2),9,WHITE)
# slide 17
s=p.slides[16];clear(s);hdr(s,'EDA Results — Multivariate Sensor Relationships',17,'MAIN PRESENTATION')
s.shapes.add_picture('outputs/charts/eda_multivariate_correlation.png',Inches(.55),Inches(1.0),width=Inches(6.1),height=Inches(4.8))
tx(s,6.95,1.0,5.65,.45,'What the heatmap adds',19,TEAL,True)
tx(s,6.95,1.52,5.45,2.3,'• Drivetrain vibration has the strongest correlation with recorded fault state (0.54).\n• Oil-particle count is also material (0.47); bearing-frequency vibration is 0.23.\n• Oil pressure is negative (−0.17); gearbox-bearing temperature is weaker (0.10).',14)
card=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(6.95),Inches(4.15),Inches(5.45),Inches(1.35));card.fill.solid();card.fill.fore_color.rgb=PALE;card.line.color.rgb=TEAL
tx(s,7.15,4.35,5.05,.95,'Implication: Fault-state evidence is distributed across signals. A multi-feature model is more appropriate than a single sensor threshold.',14,INK,True)
# slide 18
s=p.slides[17];clear(s);hdr(s,'Tuning Comparison — Recall First, Precision Trade-off',18,'APPENDIX')
tx(s,.6,.82,12,.38,'Candidates were selected from baseline validation recall. Test data did not influence tuning, selection or threshold choice.',14,INK)
data=[['Model','Baseline P / R / F1','Tuned P / R / F1','Configuration / method'],['ANN','50.28 / 68.05 / 57.83','37.67 / 92.49 / 53.54','64/24 units; BatchNorm; dropout; Adam; early stopping'],['XGBoost','71.01 / 60.83 / 65.53','69.57 / 69.10 / 69.34','6 randomized configurations; two expanding time folds']]
t=s.shapes.add_table(3,4,Inches(.55),Inches(1.35),Inches(12.2),Inches(1.7)).table
for i,row in enumerate(data):
 for j,v in enumerate(row):
  c=t.cell(i,j);c.text=v;c.fill.solid();c.fill.fore_color.rgb=NAVY if i==0 else (PALE if i%2 else WHITE)
  for q in c.text_frame.paragraphs:q.font.size=Pt(10 if i else 11);q.font.bold=i==0;q.font.color.rgb=WHITE if i==0 else INK;q.alignment=PP_ALIGN.CENTER
card=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.7),Inches(3.55),Inches(11.8),Inches(1.75));card.fill.solid();card.fill.fore_color.rgb=PALE;card.line.color.rgb=TEAL
tx(s,.95,3.83,11.3,1.2,'Selection rationale\nTuned ANN has the highest validation recall (92.49%) and therefore identifies more recorded fault observations under the declared recall-first criterion. XGBoost has higher precision and F1, making it a lower-alert-burden alternative, but its validation recall is 69.10%.',15,INK,False)
p.save('outputs/AIML_Project_1_Low_Code_Final_Presentation.pptx');print('saved',len(p.slides))
