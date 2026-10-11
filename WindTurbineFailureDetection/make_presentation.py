import json,os
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from matplotlib.backends.backend_pdf import PdfPages
import matplotlib.pyplot as plt
r=json.load(open('outputs/results.json'));out='outputs'
prs=Presentation('AIML_Project_1_Low_Code_Presentation_Template.pptx')
while len(prs.slides):
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
slides=[
('Wind Turbine Failure Detection','Aeolus Renewables | Low-Code Machine Learning Assignment\nDrivetrain fault early warning using 10-minute SCADA data'),
('Executive Summary',f"Selected model: {r['final_model']}\\nUntouched test recall: {r['test_metrics']['recall']:.2%}; only {r['test_metrics']['fn']} missed failure\\nPrecision: {r['test_metrics']['precision']:.2%}; staged review manages {r['test_metrics']['fp']} false alarms\\nRecommendation: use as maintenance-priority decision support."),
('Business Problem and Approach','Problem: detect drivetrain faults early while minimizing costly missed failures.\\nApproach: chronological 70/15/15 split; training-only preprocessing; class-aware models; validation recall for selection; untouched test evaluation.'),
('EDA Results','131,760 10-minute records from 15 turbines; fault prevalence 3.00%.\\nThermal, vibration and oil-condition signals separate normal and fault observations.\\nTurbine-specific rates support asset-level prioritization.', 'outputs/charts/eda_failure_signals.png'),
('Preprocessing','Duplicates: 0. Missing readings: gearbox oil temperature (645), generator bearing temperature (690), oil pressure (636).\\nTraining-median imputation; engineered time cycles, power utilisation and temperature rise.\\nTimestamp and turbine ID excluded. Chronological splits prevent future leakage.'),
('Baseline Models','Validation metrics: Decision Tree recall 47.93%; Random Forest 18.29%; Gradient Boosting 73.72%; XGBoost 59.67%; ANN 71.41%.\\nAccuracy alone is insufficient for this 3% fault-rate outcome.'),
('Hyperparameter Tuning','Strongest recall baselines: Gradient Boosting and ANN.\\nTuned Gradient Boosting: validation recall 99.81%, F1 60.79%.\\nTuned ANN: validation recall 93.36%, F1 67.04%.\\nSelection used validation metrics only; no test tuning.'),
('Final Model and Drivers',f"Final: tuned Gradient Boosting, selected for highest validation recall.\\nTest: accuracy {r['test_metrics']['accuracy']:.2%}, precision {r['test_metrics']['precision']:.2%}, recall {r['test_metrics']['recall']:.2%}, F1 {r['test_metrics']['f1']:.2%}.\\nConfusion matrix total: {sum(r['test_metrics'][x] for x in ['tn','fp','fn','tp'])} = test size {r['splits']['test']}.", 'outputs/charts/feature_importance.png'),
('Recommendations','1. Rank persistent alerts for operations-centre review and field dispatch.\\n2. Use vibration, temperature rise, oil and maintenance signals in technician checklists.\\n3. Resolve false alarms remotely before dispatch where possible.\\n4. Monitor missed faults, precision, lead time, downtime avoided and technician workload after deployment.')]
def page(title,body,img=None):
 s=prs.slides.add_slide(prs.slide_layouts[6]);bg=s.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(11,28,50)
 t=s.shapes.add_textbox(Inches(.65),Inches(.45),Inches(12),Inches(.65)).text_frame;t.text=title;t.paragraphs[0].font.size=Pt(30);t.paragraphs[0].font.bold=True;t.paragraphs[0].font.color.rgb=RGBColor(75,205,196)
 b=s.shapes.add_textbox(Inches(.75),Inches(1.35),Inches(6.1 if img else 11.5),Inches(5.4)).text_frame;b.text=body
 for p in b.paragraphs:p.font.size=Pt(19);p.font.color.rgb=RGBColor(238,242,247);p.space_after=Pt(14)
 if img and os.path.exists(img):s.shapes.add_picture(img,Inches(7),Inches(1.45),width=Inches(5.6))
for z in slides:page(*z)
prs.save(out+'/AIML_Project_1_Low_Code_Final_Presentation.pptx')
with PdfPages(out+'/AIML_Project_1_Low_Code_Final_Presentation.pdf') as pdf:
 for title,body,*rest in slides:
  fig=plt.figure(figsize=(13.333,7.5));fig.patch.set_facecolor('#0b1c32');plt.axis('off');plt.text(.05,.9,title,fontsize=25,color='#4bcdc4',weight='bold',transform=fig.transFigure);plt.text(.07,.72,body,fontsize=16,color='white',va='top',wrap=True,transform=fig.transFigure)
  if rest and os.path.exists(rest[0]):im=plt.imread(rest[0]);ax=fig.add_axes([.54,.12,.4,.55]);ax.imshow(im);ax.axis('off')
  pdf.savefig(fig);plt.close(fig)
print('created presentation and PDF')
